from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from uri_backend.ingestion.contracts import (
    AdapterInput,
    NormalizationResult,
    NormalizationWarning,
    NormalizedPart,
)

DOCUMENT_MEDIA_TYPES = frozenset(
    {
        "text/plain",
        "text/markdown",
        "text/x-markdown",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/pdf",
    }
)
MIN_PDF_TEXT_CHARACTERS = 20


class DocumentAdapter:
    name = "document"
    version = "normalization-v1"

    def supports(self, context: AdapterInput) -> bool:
        return context.family == "document" and context.media_type in DOCUMENT_MEDIA_TYPES

    def normalize(self, context: AdapterInput) -> NormalizationResult:
        if not self.supports(context):
            return self._result("unsupported", [], [NormalizationWarning(code="unsupported_document", message="Unsupported document family or media type.")], 0.0)
        try:
            if context.media_type in {"text/plain", "text/markdown", "text/x-markdown"}:
                return self._normalize_text(context.artifact_path)
            if context.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                return self._normalize_docx(context.artifact_path)
            return self._normalize_pdf(context.artifact_path)
        except UnicodeDecodeError as error:
            return self._result("failed", [], [NormalizationWarning(code="encoding_error", message=str(error))], 0.0)
        except Exception as error:  # noqa: BLE001 - parser failures are surfaced, never converted into text.
            return self._result("failed", [], [NormalizationWarning(code="parse_error", message=str(error))], 0.0)

    def _normalize_text(self, artifact_path: Path) -> NormalizationResult:
        content = artifact_path.read_text(encoding="utf-8", errors="strict")
        parts: list[NormalizedPart] = []
        start: int | None = None
        lines: list[str] = []
        for line_number, line in enumerate(content.splitlines(), start=1):
            if line.strip():
                if start is None:
                    start = line_number
                lines.append(line)
            elif start is not None:
                parts.append(self._part(len(parts) + 1, "paragraph", "\n".join(lines), {"line_start": start, "line_end": line_number - 1}))
                start, lines = None, []
        if start is not None:
            parts.append(self._part(len(parts) + 1, "paragraph", "\n".join(lines), {"line_start": start, "line_end": start + len(lines) - 1}))
        warnings = [] if parts else [NormalizationWarning(code="empty_document", message="Document contains no non-blank text.")]
        return self._result("normalized", parts, warnings, 1.0)

    def _normalize_docx(self, artifact_path: Path) -> NormalizationResult:
        document = Document(artifact_path)
        parts: list[NormalizedPart] = []
        for index, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if text:
                kind = "heading" if paragraph.style.name.lower().startswith("heading") else "paragraph"
                parts.append(self._part(len(parts) + 1, kind, text, {"paragraph": index}))
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = "\t".join(cell.text.strip() for cell in row.cells)
                if text.strip():
                    parts.append(self._part(len(parts) + 1, "table_row", text, {"table": table_index, "row": row_index}))
        warnings = [] if parts else [NormalizationWarning(code="empty_document", message="Document contains no non-blank text.")]
        return self._result("normalized", parts, warnings, 1.0)

    def _normalize_pdf(self, artifact_path: Path) -> NormalizationResult:
        reader = PdfReader(artifact_path)
        extracted = [(index, page.extract_text() or "") for index, page in enumerate(reader.pages, start=1)]
        text_characters = sum(len(re.sub(r"\s+", "", text)) for _, text in extracted)
        if text_characters < MIN_PDF_TEXT_CHARACTERS:
            return self._result("needs_ocr", [], [NormalizationWarning(code="no_extractable_text", message="PDF has insufficient extractable text and requires OCR.")], 0.0)
        parts = [self._part(index, "page", text.strip(), {"page": index}) for index, text in extracted if text.strip()]
        return self._result("normalized", parts, [], len(parts) / len(extracted) if extracted else 0.0)

    def _part(self, ordinal: int, kind: str, text: str, locator: dict[str, int]) -> NormalizedPart:
        return NormalizedPart(ordinal=ordinal, kind=kind, text=text, locator=locator)

    def _result(self, status: str, parts: list[NormalizedPart], warnings: list[NormalizationWarning], coverage: float) -> NormalizationResult:
        return NormalizationResult(adapter=self.name, adapter_version=self.version, status=status, parts=parts, warnings=warnings, parse_coverage=coverage)
